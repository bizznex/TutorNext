"""Create the daily GitHub issue for the 90-day Mountain Plan.

The Excel workbook is deliberately the source of truth.  This script reads the
daily plan and hard gates from it, measures progress from prior GitHub issues,
and creates the issue for the requested date.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = ROOT / "Google_90_Day_Bulletproof_Mountain_Plan.xlsx"
MARKER_RE = re.compile(r"<!-- mountain-plan:v1 day=(\d+) date=(\d{4}-\d{2}-\d{2}) -->")
CHECKBOX_RE = re.compile(r"^- \[([ xX])\] (.+)$", re.MULTILINE)

CHECK_HOUR_1 = "Hour 1 complete"
CHECK_HOUR_2 = "Hour 2 complete"
CHECK_EVIDENCE = "Required evidence linked below"
CHECK_HANDOFF = "Codex handoff reviewed"


@dataclass(frozen=True)
class DayPlan:
    day: int
    week: int
    date: date
    phase: str
    focus: str
    hour_1: str
    hour_2: str
    handoff: str
    evidence: str
    gate: str


@dataclass(frozen=True)
class Gate:
    name: str
    due_day: int
    pass_condition: str


@dataclass(frozen=True)
class Metrics:
    elapsed_days: int
    issues_found: int
    completed_hours: int
    scheduled_hours: int
    evidence_days: int
    completed_days: int
    handoff_days: int

    @property
    def pace_percent(self) -> int:
        if self.scheduled_hours == 0:
            return 100
        return round(100 * self.completed_hours / self.scheduled_hours)

    @property
    def evidence_percent(self) -> int:
        if self.elapsed_days == 0:
            return 100
        return round(100 * self.evidence_days / self.elapsed_days)

    @property
    def signal(self) -> str:
        score = min(self.pace_percent, self.evidence_percent)
        if self.elapsed_days == 0:
            return "Baseline"
        if score >= 90:
            return "On track"
        if score >= 70:
            return "Watch"
        return "Recovery required"


def _as_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def load_plan(workbook_path: Path) -> tuple[list[DayPlan], list[Gate]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        daily = workbook["Daily 2-Hour Plan"]
        plans: list[DayPlan] = []
        for row in daily.iter_rows(min_row=2, max_col=10, values_only=True):
            if row[0] is None:
                continue
            plans.append(
                DayPlan(
                    day=int(row[0]),
                    week=int(row[1]),
                    date=_as_date(row[2]),
                    phase=str(row[3]),
                    focus=str(row[4]),
                    hour_1=str(row[5]),
                    hour_2=str(row[6]),
                    handoff=str(row[7]),
                    evidence=str(row[8]),
                    gate=str(row[9]),
                )
            )

        dashboard = workbook["Executive Dashboard"]
        gates = [
            Gate(str(row[0]), int(row[1]), str(row[2]))
            for row in dashboard.iter_rows(min_row=17, max_row=23, min_col=1, max_col=3, values_only=True)
            if row[0] is not None
        ]
    finally:
        workbook.close()

    if len(plans) != 90:
        raise ValueError(f"Expected 90 daily plan rows, found {len(plans)}")
    expected_days = list(range(1, 91))
    if [item.day for item in plans] != expected_days:
        raise ValueError("Daily plan day numbers must be sequential from 1 to 90")
    for previous, current in zip(plans, plans[1:]):
        if current.date != previous.date + timedelta(days=1):
            raise ValueError("Daily plan dates must be consecutive")
    return plans, gates


class GitHubClient:
    def __init__(self, repository: str, token: str) -> None:
        self.repository = repository
        self.token = token
        self.base = f"https://api.github.com/repos/{repository}"

    def request(self, method: str, path: str, payload: dict[str, Any] | None = None) -> Any:
        body = None if payload is None else json.dumps(payload).encode("utf-8")
        request = Request(
            f"{self.base}{path}",
            data=body,
            method=method,
            headers={
                "Accept": "application/vnd.github+json",
                "Authorization": f"Bearer {self.token}",
                "X-GitHub-Api-Version": "2022-11-28",
                "User-Agent": "mountain-plan-coach",
            },
        )
        try:
            with urlopen(request, timeout=30) as response:
                return json.loads(response.read().decode("utf-8"))
        except HTTPError as error:
            details = error.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"GitHub API {method} {path} failed ({error.code}): {details}") from error

    def list_plan_issues(self) -> list[dict[str, Any]]:
        query = urlencode({"state": "all", "labels": "mountain-plan", "per_page": 100})
        return self.request("GET", f"/issues?{query}")

    def ensure_label(self, name: str, color: str, description: str) -> None:
        try:
            self.request("POST", "/labels", {"name": name, "color": color, "description": description})
        except RuntimeError as error:
            if "422" not in str(error):
                raise

    def create_issue(self, title: str, body: str, assignee: str | None) -> dict[str, Any]:
        payload: dict[str, Any] = {"title": title, "body": body, "labels": ["mountain-plan", "daily-plan"]}
        if assignee:
            payload["assignees"] = [assignee]
        return self.request("POST", "/issues", payload)


def issue_day(issue: dict[str, Any]) -> int | None:
    match = MARKER_RE.search(issue.get("body") or "")
    return int(match.group(1)) if match else None


def checked_items(body: str) -> set[str]:
    return {label.strip() for mark, label in CHECKBOX_RE.findall(body or "") if mark.lower() == "x"}


def calculate_metrics(plans: list[DayPlan], issues: list[dict[str, Any]], before_day: int) -> Metrics:
    elapsed_days = max(0, min(before_day - 1, len(plans)))
    by_day = {day: issue for issue in issues if (day := issue_day(issue)) is not None}
    completed_hours = evidence_days = completed_days = handoff_days = 0
    issues_found = 0
    for day in range(1, elapsed_days + 1):
        issue = by_day.get(day)
        if not issue:
            continue
        issues_found += 1
        checked = checked_items(issue.get("body") or "")
        hour_1 = CHECK_HOUR_1 in checked
        hour_2 = CHECK_HOUR_2 in checked
        evidence = CHECK_EVIDENCE in checked
        handoff = CHECK_HANDOFF in checked
        completed_hours += int(hour_1) + int(hour_2)
        evidence_days += int(evidence)
        handoff_days += int(handoff)
        completed_days += int(hour_1 and hour_2 and evidence)
    return Metrics(
        elapsed_days=elapsed_days,
        issues_found=issues_found,
        completed_hours=completed_hours,
        scheduled_hours=elapsed_days * 2,
        evidence_days=evidence_days,
        completed_days=completed_days,
        handoff_days=handoff_days,
    )


def next_gate(gates: list[Gate], current_day: int) -> Gate:
    return next((gate for gate in gates if gate.due_day >= current_day), gates[-1])


def build_issue_body(plan: DayPlan, gate: Gate, metrics: Metrics) -> str:
    days_to_gate = gate.due_day - plan.day
    gate_timing = "today" if days_to_gate == 0 else f"in {days_to_gate} day{'s' if days_to_gate != 1 else ''}"
    return f"""<!-- mountain-plan:v1 day={plan.day} date={plan.date.isoformat()} -->
# Day {plan.day}/90 · {plan.focus}

**Start:** 9:00 AM IST<br>
**Week / phase:** Week {plan.week} · {plan.phase}<br>
**Gate contribution:** {plan.gate}

## Your focused two hours

- [ ] {CHECK_HOUR_1} — **9:00–10:00:** {plan.hour_1}
- [ ] {CHECK_HOUR_2} — **10:00–11:00:** {plan.hour_2}
- [ ] {CHECK_EVIDENCE} — {plan.evidence}
- [ ] {CHECK_HANDOFF} — {plan.handoff}

## Performance before today

| Signal | Pace | Evidence discipline | Fully completed days |
|---|---:|---:|---:|
| **{metrics.signal}** | {metrics.completed_hours}/{metrics.scheduled_hours} hours ({metrics.pace_percent}%) | {metrics.evidence_days}/{metrics.elapsed_days} days ({metrics.evidence_percent}%) | {metrics.completed_days}/{metrics.elapsed_days} |

Missing daily issues count as zero progress so the score cannot hide skipped days.

## Next hard gate

**{gate.name} — Day {gate.due_day} ({gate_timing})**<br>
{gate.pass_condition}

## Evidence and learning

Add links to commits, documents, interview notes, recordings, or decisions here before checking the evidence box.

- Evidence:
- Blocker / learning:
- Decision made:

## End-of-session rule

Check the boxes truthfully, add evidence, and close this issue only when the day's evidence exists. The next morning's coach issue will score these checkboxes.
"""


def select_plan_for_date(plans: list[DayPlan], run_date: date) -> DayPlan | None:
    return next((plan for plan in plans if plan.date == run_date), None)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--date", dest="run_date", help="Run date in YYYY-MM-DD (defaults to Asia/Kolkata today)")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--dry-run", action="store_true", help="Print the issue instead of calling GitHub")
    return parser.parse_args(argv)


def today_in_india() -> date:
    try:
        from zoneinfo import ZoneInfo

        return datetime.now(ZoneInfo("Asia/Kolkata")).date()
    except Exception:
        return (datetime.utcnow() + timedelta(hours=5, minutes=30)).date()


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    run_date = date.fromisoformat(args.run_date) if args.run_date else today_in_india()
    plans, gates = load_plan(args.workbook)
    plan = select_plan_for_date(plans, run_date)

    if plan is None:
        print(f"{run_date} is outside the Mountain Plan ({plans[0].date} to {plans[-1].date}); no issue created.")
        return 0

    repository = os.getenv("GITHUB_REPOSITORY", "")
    token = os.getenv("GITHUB_TOKEN", "")
    issues: list[dict[str, Any]] = []
    client: GitHubClient | None = None
    if not args.dry_run:
        if not repository or not token:
            raise RuntimeError("GITHUB_REPOSITORY and GITHUB_TOKEN are required unless --dry-run is used")
        client = GitHubClient(repository, token)
        issues = client.list_plan_issues()

    metrics = calculate_metrics(plans, issues, before_day=plan.day)
    gate = next_gate(gates, plan.day)
    title = f"[Mountain Plan] Day {plan.day:02d} — {plan.focus} ({plan.date.isoformat()})"
    body = build_issue_body(plan, gate, metrics)

    if args.dry_run:
        print(title)
        print(body)
        return 0

    assert client is not None
    duplicate = next((issue for issue in issues if issue_day(issue) == plan.day), None)
    if duplicate:
        print(f"Day {plan.day} issue already exists: {duplicate['html_url']}")
        return 0

    client.ensure_label("mountain-plan", "5319E7", "90-day Mountain Plan accountability")
    client.ensure_label("daily-plan", "0E8A16", "Daily two-hour execution issue")
    assignee = os.getenv("MOUNTAIN_ASSIGNEE") or os.getenv("GITHUB_ACTOR")
    issue = client.create_issue(title, body, assignee)
    print(f"Created {issue['html_url']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
